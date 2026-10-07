# -*- coding: utf-8 -*-
"""
OCARCARE Grand Prix — build docs/data.js từ báo cáo 17h + file mapping.

Cách dùng:
    python tools/build_data.py                 # lấy file mới nhất trong input/ và mapping/
    python tools/build_data.py --report "đường_dẫn.xlsx" --date 2026-09-29

Chỉ xuất số liệu ĐÃ TỔNG HỢP theo đơn vị (không có tên KH, SĐT, biển số...).
Yêu cầu: pip install openpyxl
"""
import argparse, datetime as dt, glob, json, os, re, sys, unicodedata
from collections import defaultdict

try:
    import openpyxl
except ImportError:
    sys.exit("Thiếu thư viện openpyxl. Chạy: pip install openpyxl")

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

# ---------------------------------------------------------------- CẤU HÌNH THỂ LỆ
CONFIG = {
    "program": {
        "name": "Chào sân OCARCARE — Launch Sprint",
        "start": "2026-09-21",
        "end": "2026-10-31",
        "channel_keyword": "OPES",          # lọc cột "Danh mục Kênh bán hàng"
    },
    "early": {"tiers": [
        {"from": 1,   "to": 300, "reward": 300000, "label": "300 HĐBH đầu tiên"},
        {"from": 301, "to": 500, "reward": 200000, "label": "200 HĐBH tiếp theo"},
        {"from": 501, "to": 600, "reward": 100000, "label": "100 HĐBH cuối cùng"},
    ]},
    "individual": {"threshold": 125_000_000, "prizes": [5_000_000, 4_000_000, 3_000_000],
                   "note": "Sẽ cập nhật sau 31/10/2026"},
    "branch": {"threshold": 250_000_000, "prizes": [7_000_000, 5_000_000, 3_000_000]},
    "hub": {"threshold": 1_250_000_000, "prizes": [15_000_000, 10_000_000],
            "milestones": [
                {"value": 1_250_000_000, "label": "Điều kiện"},
                {"value": 1_825_000_000, "label": "+5 triệu"},
                {"value": 2_500_000_000, "label": "+10 triệu"}],
            "manager": {"threshold": 400_000_000,
                        "prizes": [6_000_000, 5_000_000, 4_000_000, 3_000_000],
                        "note": "Chưa phân tách được cấp phòng/nhóm Hub — công bố theo kết quả chính thức"}},
}

# ---------------------------------------------------------------- tiện ích
def norm(s):
    if s is None:
        return ""
    s = unicodedata.normalize("NFC", str(s)).replace("\n", " ")
    return re.sub(r"\s+", " ", s).strip().lower()

def to_date(v):
    if v is None or v == "":
        return None
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    s = str(v).strip()
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y %H:%M:%S", "%Y-%m-%d %H:%M:%S"):
        try:
            return dt.datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return None

def to_num(v):
    if v is None or v == "":
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).replace(",", "").replace(" ", "")
    try:
        return float(s)
    except ValueError:
        return 0.0

def latest(pattern):
    files = [f for f in glob.glob(pattern) if not os.path.basename(f).startswith("~$")]
    if not files:
        return None
    return max(files, key=os.path.getmtime)

def find_col(rows, *keys, start=0):
    """Tìm cột có tiêu đề chứa tất cả keys (đã chuẩn hoá) trong các dòng tiêu đề."""
    keys = [norm(k) for k in keys]
    for r in rows:
        for j, v in enumerate(r):
            if j < start:
                continue
            n = norm(v)
            if n and all(k in n for k in keys):
                return j
    return None

# ---------------------------------------------------------------- đọc mapping
def load_mapping(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    hdr = rows[:3]
    c_type = find_col(hdr, "cn/pgd")
    c_name = find_col(hdr, "tên cn-pgd")
    c_code = find_col(hdr, "branch code", "mã cn")
    c_pname = find_col(hdr, "tên cn quản lý")
    c_pcode = find_col(hdr, "mã cn mẹ")
    c_region = find_col(hdr, "miền")
    c_prov = find_col(hdr, "tỉnh/thành phố")
    c_pvi = find_col(hdr, "tên đơn vị pvi phụ trách")
    need = dict(type=c_type, name=c_name, code=c_code, pname=c_pname, pcode=c_pcode)
    miss = [k for k, v in need.items() if v is None]
    if miss:
        sys.exit(f"File mapping thiếu cột: {miss}")

    def g(r, c):
        return (str(r[c]).strip() if c is not None and c < len(r) and r[c] is not None else "")

    m = {}
    for r in rows[3:]:
        code = g(r, c_code).upper()
        # mã điểm bán: VN… (CN/PGD/TTTC/TTKD), HH…/HS… (House Hold – HHB)…
        if not re.fullmatch(r"[A-Z]{2}\d{6,}", code):
            continue
        m[code] = dict(code=code, type=g(r, c_type).upper(), name=g(r, c_name),
                       pname=g(r, c_pname), pcode=g(r, c_pcode).upper(),
                       region=g(r, c_region), province=g(r, c_prov), pvi=g(r, c_pvi))
    return m

def unit_of(rec, mapping):
    """Trả về (loại, mã đơn vị thi đua, tên).
    Từ 07/10/2026 (thống nhất với OPES): CN, PGD, TTTC, House Hold (HHB) là các đơn vị ĐỒNG CẤP —
    mỗi điểm bán là một đơn vị thi đua riêng, ngưỡng 250tr áp dụng cho từng điểm.
    TTKD OTO MIỀN BẮC/NAM là Hub xe (hạng mục 3)."""
    t = (rec["type"] or "").upper()
    if "TTKD" in t or "OTO" in rec["name"].upper():
        return "HUB", rec["code"], rec["name"]
    if "PGD" in t:
        kind = "PGD"
    elif "TTTC" in t:
        kind = "TTTC"
    elif "HHB" in t or rec["name"].upper().startswith("HH"):
        kind = "HHB"
    else:
        kind = "CN"
    return kind, rec["code"], rec["name"]

# ---------------------------------------------------------------- đọc báo cáo
def load_report(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = None
    for s in wb.worksheets:
        if "oto" in s.title.lower() or "nvu" in s.title.lower():
            ws = s
            break
    ws = ws or wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    hi = None
    for i, r in enumerate(rows[:40]):
        ns = [norm(v) for v in r]
        if "số đơn bảo hiểm" in ns and any("mã chi nhánh bank" in n for n in ns):
            hi = i
            break
    if hi is None:
        sys.exit("Không tìm thấy dòng tiêu đề (cần có 'Số đơn bảo hiểm' và 'Mã chi nhánh Bank').")
    hdr = rows[hi:hi + 2]
    c = dict(
        pol=find_col([rows[hi]], "số đơn bảo hiểm"),
        sdbs=find_col([rows[hi]], "số đơn sửa đổi bổ sung"),
        issued=find_col([rows[hi]], "ngày cấp"),
        channel=find_col([rows[hi]], "danh mục kênh bán hàng"),
        fee=find_col([rows[hi]], "tổng phí bh"),
        bank=find_col([rows[hi]], "mã chi nhánh bank"),
    )
    if c["fee"] is None or c["bank"] is None or c["pol"] is None or c["issued"] is None:
        sys.exit(f"Báo cáo thiếu cột bắt buộc: {c}")
    out = []
    for r in rows[hi + 3:]:
        if c["pol"] >= len(r):
            continue
        pol = r[c["pol"]]
        if not pol or not str(pol).strip() or "/" not in str(pol):
            continue
        out.append(dict(
            pol=str(pol).strip(),
            sdbs=str(r[c["sdbs"]] or "").strip() if c["sdbs"] is not None else "",
            issued=to_date(r[c["issued"]]),
            channel=str(r[c["channel"]] or "") if c["channel"] is not None else "",
            fee=to_num(r[c["fee"]]),
            bank=str(r[c["bank"]] or "").strip().upper(),
        ))
    # dòng SĐBS để trống Mã chi nhánh Bank → kế thừa từ đơn gốc
    origin = {x["pol"]: x["bank"] for x in out if x["bank"] and not x["sdbs"]}
    for x in out:
        if not x["bank"]:
            x["bank"] = origin.get(x["pol"], "")
    return out

# ---------------------------------------------------------------- tổng hợp
def rank(items):
    items.sort(key=lambda x: (-x["premium"], -x["policies"], x["name"]))
    for i, x in enumerate(items, 1):
        x["rank"] = i
    return items

def aggregate(recs, mapping, cutoff=None):
    """cutoff: chỉ tính đơn có ngày cấp < cutoff (để tính 'hôm qua')."""
    units = {}
    seen = set()
    unmapped = {"premium": 0.0, "policies": 0, "codes": set()}
    for x in recs:
        if cutoff and x["issued"] >= cutoff:
            continue
        rec = mapping.get(x["bank"])
        new_pol = not x["sdbs"] and x["pol"] not in seen
        if new_pol:
            seen.add(x["pol"])
        if not rec:
            unmapped["premium"] += x["fee"]
            unmapped["policies"] += 1 if new_pol else 0
            unmapped["codes"].add(x["bank"] or "(trống)")
            continue
        kind, uid, uname = unit_of(rec, mapping)
        u = units.get(uid)
        if not u:
            base = mapping.get(uid, rec)
            u = units[uid] = dict(id=uid, kind=kind, name=uname,
                                  region=base["region"] or rec["region"],
                                  province=base["province"] or rec["province"],
                                  pvi=base["pvi"] or rec["pvi"],
                                  premium=0.0, policies=0)
        u["premium"] += x["fee"]
        u["policies"] += 1 if new_pol else 0
    return units, unmapped, len(seen)

def aggregate_pos(recs, mapping, rdate):
    """Tổng hợp theo TỪNG điểm bán (mã CN/PGD/TTTC/TTKD) — dùng cho góc đơn vị BH PVI.
    CN và PGD hạch toán độc lập: mỗi điểm bán thuộc đơn vị BH PVI ghi trên dòng mapping của chính nó."""
    pos, seen = {}, set()
    for x in recs:
        rec = mapping.get(x["bank"])
        new_pol = not x["sdbs"] and x["pol"] not in seen
        if new_pol:
            seen.add(x["pol"])
        if not rec:
            continue
        kind, uid, uname = unit_of(rec, mapping)
        p = pos.get(rec["code"])
        if not p:
            p = pos[rec["code"]] = dict(c=rec["code"], n=rec["name"], t=rec["type"] or kind, k=kind,
                                        u=uid, un=uname, v=rec["pvi"] or "Chưa xác định",
                                        region=rec["region"], premium=0.0, policies=0,
                                        today_premium=0.0, today_policies=0)
        p["premium"] += x["fee"]
        p["policies"] += 1 if new_pol else 0
        if x["issued"] == rdate:
            p["today_premium"] += x["fee"]
            p["today_policies"] += 1 if new_pol else 0
    out = sorted(pos.values(), key=lambda z: (-z["premium"], z["n"]))
    for z in out:
        z["premium"] = round(z["premium"]); z["today_premium"] = round(z["today_premium"])
    return out

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--report")
    ap.add_argument("--mapping")
    ap.add_argument("--date", help="Ngày báo cáo YYYY-MM-DD (mặc định lấy từ tên file)")
    ap.add_argument("--out", default=os.path.join(ROOT, "docs"))
    a = ap.parse_args()

    report = a.report or latest(os.path.join(ROOT, "input", "*.xls*"))
    mfile = a.mapping or latest(os.path.join(ROOT, "mapping", "*.xls*"))
    if not report:
        sys.exit("Không có file báo cáo trong thư mục input/")
    if not mfile:
        sys.exit("Không có file mapping trong thư mục mapping/")
    print(f"• Báo cáo : {os.path.basename(report)}")
    print(f"• Mapping : {os.path.basename(mfile)}")

    mapping = load_mapping(mfile)
    recs_all = load_report(report)

    P = CONFIG["program"]
    start, end = dt.date.fromisoformat(P["start"]), dt.date.fromisoformat(P["end"])
    kw = P["channel_keyword"].lower()
    recs = [x for x in recs_all if x["issued"] and start <= x["issued"] <= end
            and (not x["channel"] or kw in x["channel"].lower())]
    print(f"• Dòng dữ liệu: {len(recs_all)} | trong kỳ & kênh {P['channel_keyword']}: {len(recs)}")

    # ngày báo cáo
    rdate = None
    if a.date:
        rdate = dt.date.fromisoformat(a.date)
    else:
        # nhận ngày trong tên file: 30092026 | 30.09.2026 | 30-09-2026 | 30_09_2026 | 2026-09-30
        name = os.path.basename(report)
        pats = [(r"(?<!\d)(\d{1,2})[._\- ](\d{1,2})[._\- ](20\d{2})(?!\d)", "dmy"),
                (r"(?<!\d)(20\d{2})[._\-](\d{1,2})[._\-](\d{1,2})(?!\d)", "ymd"),
                (r"(?<!\d)(\d{2})(\d{2})(20\d{2})(?!\d)", "dmy")]
        for pat, order in pats:
            for mt in re.finditer(pat, name):
                a1, a2, a3 = (int(g) for g in mt.groups())
                try:
                    cand = dt.date(a3, a2, a1) if order == "dmy" else dt.date(a1, a2, a3)
                except ValueError:
                    continue
                if start <= cand <= end + dt.timedelta(days=7):
                    rdate = cand
                    break
            if rdate:
                break
    if not rdate:
        rdate = max((x["issued"] for x in recs), default=dt.date.today())
        print(f"  ℹ Tên file không có ngày → lấy ngày cấp mới nhất trong dữ liệu: {rdate:%d/%m/%Y}")

    mh = re.search(r"(?<!\d)(\d{1,2})\s*[hH](\d{2})?(?![a-zA-Z])", os.path.basename(report))
    report_time = f"{int(mh.group(1)):02d}:{mh.group(2) or '00'}" if mh and int(mh.group(1)) < 24 else "17:00"

    now_u, unmapped, total_pol = aggregate(recs, mapping)
    prev_u, _, prev_pol = aggregate(recs, mapping, cutoff=rdate)

    def board(kind_filter, thr):
        cur = rank([dict(u) for u in now_u.values() if u["kind"] in kind_filter])
        prv = {u["id"]: u for u in rank([dict(u) for u in prev_u.values() if u["kind"] in kind_filter])}
        for u in cur:
            p = prv.get(u["id"])
            u["prev_rank"] = p["rank"] if p else None
            u["today_premium"] = round(u["premium"] - (p["premium"] if p else 0))
            u["today_policies"] = u["policies"] - (p["policies"] if p else 0)
            u["qualified"] = u["premium"] >= thr
            u["crossed_today"] = u["qualified"] and not (p and p["premium"] >= thr)
            u["premium"] = round(u["premium"])
        return cur

    wk_start = rdate - dt.timedelta(days=6)
    wk_u, _, _ = aggregate(recs, mapping, cutoff=wk_start)          # trạng thái cuối ngày (rdate-7)
    unit_days = defaultdict(set); unit_first = {}
    for x in recs:
        rec = mapping.get(x["bank"])
        if not rec:
            continue
        uid = unit_of(rec, mapping)[1]
        if x["issued"] >= wk_start:
            unit_days[uid].add(x["issued"])
        if uid not in unit_first or x["issued"] < unit_first[uid]:
            unit_first[uid] = x["issued"]

    def add_week(items, kinds):
        wk = {u["id"]: u for u in rank([dict(u) for u in wk_u.values() if u["kind"] in kinds])}
        for u in items:
            w = wk.get(u["id"])
            u["rank7"] = w["rank"] if w else None
            u["week_premium"] = round(u["premium"] - (w["premium"] if w else 0))
            u["week_policies"] = u["policies"] - (w["policies"] if w else 0)
            u["active_days7"] = len(unit_days.get(u["id"], ()))
            f = unit_first.get(u["id"])
            u["first_date"] = f.isoformat() if f else None
        return items

    branches = board({"CN", "PGD", "TTTC", "HHB"}, CONFIG["branch"]["threshold"])
    hubs = board({"HUB"}, CONFIG["hub"]["threshold"])
    # luôn hiển thị đủ 2 Hub kể cả khi chưa có đơn
    for code, rec in mapping.items():
        if unit_of(rec, mapping)[0] == "HUB" and code not in {h["id"] for h in hubs}:
            hubs.append(dict(id=code, kind="HUB", name=rec["name"], region=rec["region"],
                             province=rec["province"], pvi=rec["pvi"], premium=0, policies=0,
                             prev_rank=None, today_premium=0, today_policies=0,
                             qualified=False, crossed_today=False))
    rank(hubs)
    add_week(branches, {"CN", "PGD", "TTTC", "HHB"}); add_week(hubs, {"HUB"})

    # tra cứu: mọi CN/PGD/TTTC/TTKD → đơn vị thi đua
    lookup = []
    for code, rec in mapping.items():
        kind, uid, uname = unit_of(rec, mapping)
        lookup.append(dict(c=code, v=rec["pvi"], n=rec["name"], t=rec["type"] or kind, u=uid, un=uname, k=kind,
                           p=rec["province"]))

    # luỹ kế theo ngày
    daily = []
    d = start
    seen = set()
    cum_prem, cum_pol = 0.0, 0
    by_day = defaultdict(list)
    for x in recs:
        by_day[x["issued"]].append(x)
    while d <= min(rdate, end):
        for x in by_day.get(d, []):
            cum_prem += x["fee"]
            if not x["sdbs"] and x["pol"] not in seen:
                seen.add(x["pol"])
                cum_pol += 1
        daily.append(dict(date=d.isoformat(), policies=cum_pol, premium=round(cum_prem)))
        d += dt.timedelta(days=1)

    def group(key):
        g = defaultdict(lambda: dict(premium=0.0, policies=0))
        for u in now_u.values():
            k = u[key] or "Khác"
            g[k]["premium"] += u["premium"]
            g[k]["policies"] += u["policies"]
        return sorted([dict(name=k, premium=round(v["premium"]), policies=v["policies"])
                       for k, v in g.items()], key=lambda x: -x["premium"])

    pos = aggregate_pos(recs, mapping, rdate)
    pv = {}
    for rec in mapping.values():
        v = rec["pvi"] or "Chưa xác định"
        pv.setdefault(v, dict(name=v, premium=0, policies=0, today_premium=0, today_policies=0,
                              points_total=0, points_active=0))
        pv[v]["points_total"] += 1
    for z in pos:
        g = pv[z["v"]]
        for k in ("premium", "policies", "today_premium", "today_policies"):
            g[k] += z[k]
        g["points_active"] += 1
    pvi_units = sorted(pv.values(), key=lambda g: (-g["premium"], -g["policies"], g["name"]))
    for i, g in enumerate(pvi_units, 1):
        g["rank"] = i

    total_prem = sum(x["fee"] for x in recs)
    today_recs = [x for x in recs if x["issued"] == rdate]
    data = dict(
        meta=dict(
            report_date=rdate.isoformat(), report_time=report_time, model="diem-ban",
            generated_at=dt.datetime.now().strftime("%Y-%m-%d %H:%M"),
            source=os.path.basename(report),
            total_policies=total_pol, total_premium=round(total_prem),
            today_policies=total_pol - prev_pol,
            today_premium=round(sum(x["fee"] for x in today_recs)),
            active_units=len([b for b in branches if b["policies"] > 0]),
            unmapped=dict(premium=round(unmapped["premium"]), policies=unmapped["policies"]),
        ),
        config=CONFIG,
        branches=branches, hubs=hubs, lookup=lookup, daily=daily,
        regions=group("region"), pvi_units=pvi_units, pos=pos,
    )

    # ---- kiểm tra dữ liệu
    warns = []
    cnt = defaultdict(int)
    for x in recs:
        if not x["sdbs"]:
            cnt[x["pol"]] += 1
    dups = [k for k, v in cnt.items() if v > 1]
    if dups:
        warns.append(f"{len(dups)} số HĐBH bị lặp (không phải SĐBS): {dups[:5]}")
    neg = [x for x in recs if x["fee"] < 0]
    if neg:
        warns.append(f"{len(neg)} dòng phí âm (huỷ/giảm phí), tổng {sum(x['fee'] for x in neg):,.0f}đ")
    fut = [x for x in recs if x["issued"] > rdate]
    if fut:
        warns.append(f"{len(fut)} dòng có ngày cấp SAU ngày báo cáo {rdate:%d/%m} — kiểm tra lại tên file/ngày")
    outp = [x for x in recs_all if x["issued"] and not (start <= x["issued"] <= end)]
    if outp:
        warns.append(f"{len(outp)} dòng có ngày cấp ngoài kỳ chương trình (đã loại)")
    hdir = os.path.join(a.out, "history")
    olds = sorted(f for f in glob.glob(os.path.join(hdir, "data-*.json"))
                  if os.path.basename(f) < f"data-{rdate.isoformat()}.json")
    if olds:
        try:
            with open(olds[-1], encoding="utf-8") as f:
                prevh = json.load(f)
            same_model = prevh.get("model") == "diem-ban"
            pm = {} if not same_model else {b["id"]: b for b in prevh.get("branches", []) + prevh.get("hubs", [])}
            down = [(b["name"], pm[b["id"]]["premium"], b["premium"]) for b in branches + hubs
                    if b["id"] in pm and b["premium"] < pm[b["id"]]["premium"] - 1]
            for n, o, nw in down:
                warns.append(f"Doanh thu {n} GIẢM so với {os.path.basename(olds[-1])[5:15]}: {o:,.0f} → {nw:,.0f}đ")
            if total_pol < prevh.get("total_policies", 0):
                warns.append(f"Tổng HĐBH giảm: {prevh.get('total_policies')} → {total_pol}")
        except Exception as e:
            warns.append(f"Không đọc được lịch sử để so sánh: {e}")

    # ---- tin nhắn Zalo
    def tr(v):
        return f"{v/1e9:,.2f} tỷ".replace(",", "X").replace(".", ",").replace("X", ".") if v >= 1e9 else \
               f"{v/1e6:,.1f} tr".replace(",", "X").replace(".", ",").replace("X", ".")
    BASE = "https://pvi-aviationdivision.github.io/OCARCARE_SALECONTEST2026/"
    tiers = CONFIG["early"]["tiers"]
    cur_t = next((t for t in tiers if total_pol < t["to"]), None)
    days_left = max(0, (end - rdate).days)
    L = [f"🏁 CHÀO SÂN OCARCARE – cập nhật {report_time.replace(':00','h')} ngày {rdate:%d/%m}",
         f"• Tổng: {total_pol} HĐBH | {tr(total_prem)} (chưa VAT)" + (f" | hôm nay +{data['meta']['today_policies']} HĐ" if data['meta']['today_policies'] else ""),
         (f"• Còn {cur_t['to'] - total_pol} suất thưởng {cur_t['reward']//1000}K/HĐBH" if cur_t else "• Đã đủ 600 HĐBH được thưởng"),
         f"• Còn {days_left} ngày đến vạch đích 31/10", "",
         "🏆 Top 3 điểm bán (CN/PGD/TTTC):"]
    for b in branches[:3]:
        L.append(f"{b['rank']}. {b['name']} – {tr(b['premium'])} ({b['policies']} HĐ)")
    ups = sorted([b for b in branches if b.get("rank7") and b["rank7"] - b["rank"] >= 2],
                 key=lambda b: b["rank"] - b["rank7"])[:2]
    news = [b for b in branches if b.get("first_date") and b["first_date"] >= (rdate - dt.timedelta(days=2)).isoformat()][:4]
    if ups or news:
        L.append("")
    if ups:
        L.append("🚀 Bứt phá tuần: " + "; ".join(f"{b['name']} ▲{b['rank7']-b['rank']} bậc" for b in ups))
    if news:
        L.append("🆕 Mới lên bảng: " + ", ".join(b["name"] for b in news))
    L += ["", "🚙 Hub xe: " + " – ".join(f"{('Miền Bắc' if 'BAC' in h['name'].upper() else 'Miền Nam')} {tr(h['premium'])}" for h in hubs),
          "", f"👉 Xem chi tiết & tìm CN/PGD của mình: {BASE}",
          "(Số liệu tạm tính, kết quả chính thức theo Thể lệ)"]
    zalo = "\n".join(L)
    with open(os.path.join(ROOT, "tin_nhan_zalo.txt"), "w", encoding="utf-8") as f:
        f.write(zalo)

    os.makedirs(os.path.join(a.out, "history"), exist_ok=True)
    js = "window.RACE_DATA = " + json.dumps(data, ensure_ascii=False, separators=(",", ":")) + ";\n"
    with open(os.path.join(a.out, "data.js"), "w", encoding="utf-8") as f:
        f.write(js)
    # chống cache: gắn mã phiên bản vào data.js trong index.html → người xem luôn nhận số mới
    idx = os.path.join(a.out, "index.html")
    if os.path.exists(idx):
        with open(idx, encoding="utf-8") as f:
            html = f.read()
        ver = dt.datetime.now().strftime("%Y%m%d%H%M%S")
        html2 = re.sub(r'src="data\.js(?:\?v=[^"]*)?"', f'src="data.js?v={ver}"', html)
        if html2 != html:
            with open(idx, "w", encoding="utf-8", newline="") as f:
                f.write(html2)
    with open(os.path.join(a.out, "history", f"data-{rdate.isoformat()}.json"), "w", encoding="utf-8") as f:
        json.dump({**data["meta"], **{"branches": [{k: b[k] for k in ("id", "name", "premium", "policies", "rank")} for b in branches],
                                  "hubs": [{k: b[k] for k in ("id", "name", "premium", "policies", "rank")} for b in hubs]}},
                  f, ensure_ascii=False, indent=1)

    # ---- báo cáo kiểm tra
    print(f"• Ngày báo cáo: {rdate:%d/%m/%Y}")
    print(f"• Tổng: {total_pol} HĐBH | {total_prem:,.0f}đ (chưa VAT) | hôm nay +{data['meta']['today_policies']} HĐ")
    print(f"• Điểm bán CN/PGD/TTTC/HHB có đơn: {len(branches)} | Hub: " +
          ", ".join(f"{h['name']} {h['premium']:,.0f}đ" for h in hubs))
    chk = sum(b["premium"] for b in branches) + sum(h["premium"] for h in hubs) + unmapped["premium"]
    if abs(chk - total_prem) > 5:
        print(f"  ⚠ Lệch tổng: {chk:,.0f} vs {total_prem:,.0f}")
    if unmapped["premium"] or unmapped["policies"]:
        print(f"  ⚠ {unmapped['policies']} HĐ có Mã chi nhánh Bank chưa có trong mapping: {sorted(unmapped['codes'])}")
    for w in warns:
        print(f"  ⚠ {w}")
    print("✔ Đã ghi docs/data.js")
    print("✔ Đã soạn tin nhắn Zalo: tin_nhan_zalo.txt")

if __name__ == "__main__":
    main()
